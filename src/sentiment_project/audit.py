from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from .common import ROOT, digest, now_utc, read_json, write_json


def audit_workbook(config=None, root=ROOT):
    config = config or read_json(root / "configs/project.json")
    path = root / config["workbook"]
    cached = load_workbook(path, read_only=False, data_only=True)
    formulas = load_workbook(path, read_only=False, data_only=False)
    sheet = cached["Sheet1"]
    fsheet = formulas["Sheet1"]
    as_of = pd.Timestamp(config["as_of"])
    quarantine = {(q["series_id"], q["period"]): q["reason"] for q in config["quarantine"]}
    rows, summary, issues = [], [], []
    for spec in config["series"]:
        dc, vc, sid = spec["date_column"], spec["value_column"], spec["id"]
        if sheet.cell(1, dc).value != spec["date_header"] or sheet.cell(1, vc).value != spec["value_header"]:
            raise ValueError(f"Workbook schema changed for {sid}; inspect before importing")
        group = []
        for rn in range(2, sheet.max_row + 1):
            d, v = sheet.cell(rn, dc).value, sheet.cell(rn, vc).value
            if d is None and v is None:
                continue
            if not isinstance(d, datetime) or not isinstance(v, (int, float)):
                issues.append({"series_id": sid, "row": rn, "issue": "invalid_date_or_value"})
                continue
            if fsheet.cell(rn, vc).data_type == "f":
                raise ValueError("Formula cells require recalculation verification before import")
            period = pd.Timestamp(d).normalize()
            expected_end = period.to_period(spec["frequency"]).end_time.normalize()
            if period != expected_end:
                raise ValueError(f"Non-period-end date: {sid} {period}")
            status = "source_unverified"
            reason = "No source URL, publication date or historical vintage in workbook"
            if period > as_of:
                status = "incomplete_or_future_period"
                reason = "Period not complete at cutoff; preliminary/final status unknown"
            key = (sid, period.date().isoformat())
            if key in quarantine:
                status, reason = "quarantined", quarantine[key]
            row = {
                "series_id": sid, "period_end": period.date().isoformat(), "value": float(v),
                "frequency": spec["frequency"], "unit": spec["unit"],
                "source_file": path.name, "source_sheet": sheet.title,
                "source_cell": sheet.cell(rn, vc).coordinate,
                "release_date": "", "vintage_date": "", "source_url": "",
                "status": status, "note": reason,
            }
            rows.append(row)
            group.append(row)
        if not group:
            raise ValueError(f"No observations: {sid}")
        periods = pd.PeriodIndex([r["period_end"] for r in group], freq=spec["frequency"])
        if periods.duplicated().any():
            raise ValueError(f"Duplicate periods in {sid}")
        missing = pd.period_range(periods.min(), periods.max(), freq=spec["frequency"]).difference(periods)
        summary.append({
            "series_id": sid, "frequency": spec["frequency"], "observations": len(group),
            "first_period": min(r["period_end"] for r in group),
            "last_period": max(r["period_end"] for r in group),
            "missing_periods": [str(x) for x in missing],
            "quarantined": sum(r["status"] == "quarantined" for r in group),
            "incomplete_or_future": sum(r["status"] == "incomplete_or_future_period" for r in group),
        })
    all_cells = [c for row in fsheet for c in row if c.value is not None]
    result = {
        "as_of": config["as_of"], "generated_at": now_utc(), "file": path.name,
        "sha256": digest(path.read_bytes()), "sheets": cached.sheetnames,
        "formula_cells": sum(c.data_type == "f" for c in all_cells),
        "hyperlinks": sum(bool(c.hyperlink) for c in all_cells),
        "comments": sum(bool(c.comment) for c in all_cells),
        "series": summary, "invalid_rows": issues,
        "strict_realtime_ready": False,
        "missing_required": ["publication dates", "historical vintages", "original PCE index", "Michigan current/expectations subindices", "household controls", "media corpus"],
    }
    cached.close()
    formulas.close()
    out = root / "data/processed"
    out.mkdir(parents=True, exist_ok=True)
    frame = pd.DataFrame(rows).sort_values(["series_id", "period_end"])
    frame.to_csv(out / "workbook_observations.csv", index=False)
    frame[frame.status != "source_unverified"].to_csv(out / "quarantined_or_incomplete.csv", index=False)
    # Monthly columns join by their OWN dates; quarterly GDP is intentionally separate.
    monthly = frame[frame.frequency == "M"].pivot(index="period_end", columns="series_id", values="value")
    monthly = monthly.reindex(pd.date_range(monthly.index.min(), monthly.index.max(), freq="ME").strftime("%Y-%m-%d"))
    monthly.index.name = "period_end"
    monthly.to_csv(out / "monthly_observed.csv")
    write_json(root / "reports/data_audit.json", result)
    lines = ["# 数据审计（第一阶段）", "", f"研究截止日：{config['as_of']}。原工作簿未修改。", "",
             "**结论：可用于探索性分析，不能直接作为实时预测回测的数据集。**", "",
             "四组序列各有日期列，已按日期独立导出与合并；GDP 保持季度频率。", "",
             "| 序列 | 条数 | 起始 | 结束 | 缺失时期 |", "|---|---:|---|---|---|"]
    for s in summary:
        lines.append(f"| {s['series_id']} | {s['observations']} | {s['first_period']} | {s['last_period']} | {', '.join(s['missing_periods']) or '无'} |")
    lines += ["", "## 需要处理的问题", "",
              "- 工作簿没有单元格公式、来源链接或批注；没有发布日期、历史版本和初值/终值标记。所有数值的正式来源仍需验证。",
              "- 2026Q3 GDP 已隔离：截至研究截止日，其 BEA 首次发布尚未发生。不能把表内值当作当时可得的官方数据。",
              "- 2026 年 10 月信心为未完成月份，可能是初值或其他口径；标为待核实，不作为已公布终值使用。",
              "- PCE 仅有同比结果，缺少价格指数，暂时不能独立复算或构造累计价格水平。",
              "- 缺失月份保持缺失；GDP 不插值。不同序列不能按行号拼接。",
              "- 后续示范回测仅检查代码与模型行为，使用未验证历史修订值及假设发布时间，不能视为最终研究证据。", "",
              "## 已核实的外部依据", "",
              "- [BEA 发布日程](https://www.bea.gov/news/schedule)：2026Q3 GDP advance 计划于 2026-10-29 发布（2026-10-16 查阅）。",
              "- [FRED GDP](https://fred.stlouisfed.org/series/A191RL1Q225SBEA)：截至查阅时最新为 2026Q2，2.2%。",
              "- [FRED U-3](https://fred.stlouisfed.org/series/UNRATE)：截至查阅时 2026-09 为 4.2%。与工作簿最新对应值一致，但不证明全历史一致。", "",
              "## 复现", "", "运行 `python run.py audit` 重新生成；CSV 含原单元格位置、隔离原因和空白的发布日期字段。"]
    (root / "reports/data_audit.zh-CN.md").write_text("\n".join(lines) + "\n", encoding="utf-8")
    return result

