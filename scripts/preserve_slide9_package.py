"""Retain the source package while applying Artifact Tool's targeted slide-9 edits."""
import copy,io,sys,zipfile,xml.etree.ElementTree as E
from pathlib import Path
from openpyxl import load_workbook
source,candidate=map(Path,sys.argv[1:3])
NS={'c':'http://schemas.openxmlformats.org/drawingml/2006/chart','a':'http://schemas.openxmlformats.org/drawingml/2006/main','p':'http://schemas.openxmlformats.org/presentationml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
for prefix,uri in NS.items():E.register_namespace(prefix,uri)
with zipfile.ZipFile(source) as original,zipfile.ZipFile(candidate) as authored:
 blobs={n:original.read(n) for n in original.namelist()}
 slide_name='ppt/slides/slide9.xml';oldslide=E.fromstring(blobs[slide_name]);newslide=E.fromstring(authored.read(slide_name))
 oldref=oldslide.find('.//c:chart',NS);newref=newslide.find('.//c:chart',NS);assert oldref is not None and newref is not None
 newref.set('{'+NS['r']+'}id',oldref.attrib['{'+NS['r']+'}id']);blobs[slide_name]=E.tostring(newslide,encoding='utf-8',xml_declaration=True)
 chart_name='ppt/slides/charts/chart8.xml';oldchart=E.fromstring(blobs[chart_name]);newchart=E.fromstring(authored.read(chart_name))
 external=oldchart.find('c:externalData',NS);assert external is not None;newchart.append(copy.deepcopy(external))
 oldnumbers=[e.text for e in oldchart.findall('.//c:ser/c:val//c:v',NS)];newnumbers=[e.text for e in newchart.findall('.//c:ser/c:val//c:v',NS)];assert oldnumbers==newnumbers,(oldnumbers,newnumbers)
 blobs[chart_name]=E.tostring(newchart,encoding='utf-8',xml_declaration=True)
 name='ppt/embeddings/chart-data-snapshot-008.xlsx';wb=load_workbook(io.BytesIO(blobs[name]));ws=wb.worksheets[0]
 labels=[v.text for v in newchart.findall('.//c:ser/c:cat//c:strCache/c:pt/c:v',NS)];assert len(labels)==3
 for row,label in enumerate(labels,2):ws.cell(row,1,label)
 stream=io.BytesIO();wb.save(stream);blobs[name]=stream.getvalue()
 assert len(blobs)==len(original.namelist())
 temporary=candidate.with_name('preserved.pptx')
 with zipfile.ZipFile(temporary,'w',compression=zipfile.ZIP_DEFLATED) as z:
  for n,b in blobs.items():z.writestr(n,b)
temporary.replace(candidate)
print('Preserved every source package part except slide9, chart8 and its category-label workbook.')
